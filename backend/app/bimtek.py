"""Penanda sekolah peserta bimtek percepatan laporan yang laporan 50%-nya belum direviu.

Sumber: file Excel dari koordinator (mis. Hitung_Hari_Lapor_50_Data_28_Okt_dengan_Petugas.xlsx).
- Sheet data sekolah ("DB"): baris per sekolah dengan kolom NPSN, Fasilitator, Hotel (= peserta bimtek),
  "Hasil Reviu 50% (<tanggal>) Adj" dan "Hari terakhir submit".
- Sheet ringkasan ("Data <tanggal>", pivot) bila ada: daftar fasil + hari terakhir submit yang dipilih koordinator.
  Sekolah dicocokkan lewat (fasil, hari terakhir submit) karena pivot tidak memuat NPSN.

Yang disimpan hanya NPSN, hotel, dan tanggal submit (tanpa nama atau nomor), di pengaturan lokal.
Penanda hilang sendiri begitu LK mencatat sekolahnya sudah direviu.
"""

import io
import re
from datetime import date, datetime
from typing import Any

from . import db, schedule

KUNCI = "bimtek"


class BimtekError(Exception):
    pass


def _teks(v: Any) -> str:
    return "" if v is None else str(v).strip()


def _npsn(v: Any) -> str:
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return _teks(v)


def _tanggal(v: Any) -> str:
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    return ""


def _kolom(header: list[str], cocok) -> int | None:
    """Indeks kolom TERAKHIR yang cocok (kolom reviu terbaru ada paling kanan)."""
    idx = [i for i, h in enumerate(header) if cocok(h.lower())]
    return idx[-1] if idx else None


def _sama_hari(a: Any, b: Any) -> bool:
    if a is None or b is None:
        return a is None and b is None
    try:
        return abs(float(a) - float(b)) < 1e-6
    except (TypeError, ValueError):
        return False


def baca(data: bytes, nama_file: str) -> dict[str, Any]:
    if not nama_file.lower().endswith((".xlsx", ".xlsm")):
        raise BimtekError("File harus Excel (.xlsx).")
    import openpyxl

    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:
        raise BimtekError(f"File Excel tidak bisa dibuka: {e}") from e

    # Sheet data sekolah: baris pertama memuat NPSN, Fasilitator, Hotel
    db_ws, header = None, []
    for ws in wb.worksheets:
        baris = next(ws.iter_rows(max_row=1, values_only=True), ())
        h = [_teks(x) for x in baris]
        low = [x.lower() for x in h]
        if "npsn" in low and "fasilitator" in low and "hotel" in low:
            db_ws, header = ws, h
            break
    if db_ws is None:
        raise BimtekError("Sheet data sekolah tidak ketemu: butuh kolom NPSN, Fasilitator, dan Hotel di baris pertama.")

    low = [x.lower() for x in header]
    i_npsn, i_fasil, i_hotel = low.index("npsn"), low.index("fasilitator"), low.index("hotel")
    i_reviu = _kolom(header, lambda h: h.startswith("hasil reviu 50%") and h.endswith("adj"))
    if i_reviu is None:
        i_reviu = _kolom(header, lambda h: h.startswith("hasil reviu 50%"))
    if i_reviu is None:
        raise BimtekError("Kolom \"Hasil Reviu 50% (...)\" tidak ketemu.")
    i_hari = _kolom(header, lambda h: h == "hari terakhir submit")
    i_submit = _kolom(header, lambda h: h.startswith("tanggal submit"))
    m = re.search(r"\(([^)]+)\)", header[i_reviu])
    per_tanggal = m.group(1) if m else ""

    calon = []
    for r in db_ws.iter_rows(min_row=2, values_only=True):
        if len(r) <= i_reviu or not _teks(r[i_hotel]) or _teks(r[i_reviu]).lower() != "belum":
            continue
        npsn = _npsn(r[i_npsn])
        if npsn:
            calon.append({"npsn": npsn, "fasil": _teks(r[i_fasil]), "hotel": _teks(r[i_hotel]),
                          "hari": r[i_hari] if i_hari is not None else None,
                          "submit": _tanggal(r[i_submit]) if i_submit is not None else ""})

    # Pivot pilihan koordinator (opsional): kolom Fasilitator + Hari terakhir submit + Count of NPSN
    pilih = None
    for ws in wb.worksheets:
        if ws is db_ws:
            continue
        rows = ws.iter_rows(max_row=8000, values_only=True)
        h_pivot = None
        for r in rows:
            low_r = [_teks(x).lower() for x in r]
            if h_pivot is None:
                if "fasilitator" in low_r and "hari terakhir submit" in low_r:
                    h_pivot = (low_r.index("fasilitator"), low_r.index("hari terakhir submit"))
                    pilih = []
                continue
            fasil = _teks(r[h_pivot[0]])
            if fasil and fasil.lower() != "grand total":
                pilih.append((fasil, r[h_pivot[1]]))
        if pilih is not None:
            break

    if pilih is not None:
        hasil = {}
        for fasil, hari in pilih:
            for c in calon:
                if c["fasil"] == fasil and _sama_hari(c["hari"], hari):
                    hasil[c["npsn"]] = c
        terpilih = list(hasil.values())
    else:
        terpilih = calon
    if not terpilih:
        raise BimtekError("Tidak ada sekolah peserta bimtek yang belum direviu di file ini.")

    return {
        "file": nama_file,
        "per_tanggal": per_tanggal,
        "diunggah": schedule.now().isoformat(timespec="seconds"),
        "sekolah": {c["npsn"]: {"hotel": c["hotel"], "submit": c["submit"]} for c in terpilih},
    }


def simpan(info: dict[str, Any]) -> None:
    db.set_setting(KUNCI, info)


def hapus() -> None:
    db.delete_setting(KUNCI)


def info() -> dict[str, Any] | None:
    return db.get_setting(KUNCI)


def peta() -> dict[str, dict[str, str]]:
    return (info() or {}).get("sekolah", {})


def tanda(s, data: dict[str, dict[str, str]] | None = None) -> dict[str, str] | None:
    """Penanda untuk sekolah LK `s`: hanya selama LK belum mencatat reviu."""
    data = peta() if data is None else data
    if s.reviu == "Ya" or s.prioritas.strip().lower() == "selesai":
        return None
    return data.get(s.npsn.strip())
