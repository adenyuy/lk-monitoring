import io
from datetime import datetime

import openpyxl
import pytest

from app import bimtek, sheets

HEADER = ["Nama Sekolah", "NPSN", "Fasilitator", "Hotel", "Hasil Reviu 50% (15 Okt 2026) Adj",
          "Hasil Reviu 50% (28 Okt 2026) Adj", "Tanggal Submit per 28 Okt", "Hari terakhir submit"]
DB = [
    ["SD A", 111, "Fasil X", "Hotel 1", "Belum", "Belum", datetime(2026, 10, 16, 8, 0), 12.5],
    ["SD B", 222, "Fasil X", "Hotel 1", "Belum", "Diterima", datetime(2026, 10, 10), 18.0],  # sudah direviu
    ["SD C", 333, "Fasil Y", None, "Belum", "Belum", None, None],                           # bukan peserta
    ["SD D", 444, "Fasil Y", "Hotel 2", "Belum", "Belum", datetime(2026, 10, 20), 8.25],
    ["SD E", 555, "Fasil Z", "Hotel 3", "Belum", "Belum", datetime(2026, 10, 1), 27.0],     # tidak dipilih di pivot
]


def _xlsx(pivot: bool = True) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "DB"
    ws.append(HEADER)
    for r in DB:
        ws.append(r)
    if pivot:
        p = wb.create_sheet("Data 28 Okt")
        p.append(["Hotel", "(Multiple Items)"])
        p.append([])
        p.append(["Hari terakhir submit", "Fasilitator", "Count of NPSN"])
        p.append([12.5, "Fasil X", 1])
        p.append([8.25, "Fasil Y", 1])
        p.append(["Grand Total", None, 2])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_baca_pakai_pivot():
    r = bimtek.baca(_xlsx(), "data.xlsx")
    assert r["per_tanggal"] == "28 Okt 2026"
    assert r["sekolah"] == {"111": {"hotel": "Hotel 1", "submit": "2026-10-16"},
                            "444": {"hotel": "Hotel 2", "submit": "2026-10-20"}}


def test_baca_tanpa_pivot_ambil_semua_peserta_belum_reviu():
    assert set(bimtek.baca(_xlsx(pivot=False), "data.xlsx")["sekolah"]) == {"111", "444", "555"}


def test_tolak_file_aneh():
    with pytest.raises(bimtek.BimtekError):
        bimtek.baca(b"a,b", "x.csv")
    wb = openpyxl.Workbook()
    wb.active.append(["NPSN", "Nama"])
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(bimtek.BimtekError):
        bimtek.baca(buf.getvalue(), "x.xlsx")


def test_tanda_hilang_setelah_direviu():
    s = sheets.Sekolah(baris=3, petugas="P", npsn="111", nama="SD A", kab_kota="", fasil="Fasil X", nomor_fasil="",
                       prioritas="Prioritas 3", progres=60.0, submit="Ya", reviu="Tidak")
    data = {"111": {"hotel": "Hotel 1", "submit": "2026-10-16"}}
    assert bimtek.tanda(s, data) == data["111"]
    s.reviu = "Ya"
    assert bimtek.tanda(s, data) is None


def test_api_unggah_dan_hapus():
    from fastapi.testclient import TestClient

    from app.main import app
    from app.routers import bimtek as rute

    app.dependency_overrides[rute.wa_blast] = lambda: None
    c = TestClient(app)
    try:
        r = c.post("/api/bimtek", files={"file": ("d.xlsx", _xlsx(), "application/octet-stream")})
        assert r.status_code == 200 and r.json()["sekolah"] == 2 and r.json()["per_tanggal"] == "28 Okt 2026"
        assert set(bimtek.peta()) == {"111", "444"}
        assert c.post("/api/bimtek", files={"file": ("d.csv", b"x", "text/csv")}).status_code == 400
        assert c.delete("/api/bimtek").json() == {"ada": False} and bimtek.peta() == {}
    finally:
        app.dependency_overrides.clear()
